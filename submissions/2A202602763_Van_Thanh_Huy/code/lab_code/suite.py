"""Experiment orchestration; validation-only selection followed by a locked test phase."""
import argparse
import gc
import json
import shutil
from dataclasses import asdict, replace
from pathlib import Path
import numpy as np
import pandas as pd
import torch
from PIL import Image
from torchvision import transforms as T
from . import dataset, model as models, inference as inf, benchmark
from .train import Config, run, run_dir, set_seed, write_json
from eval import compute_metrics, save_predictions, read_pred

METHODS = ['single', 'hflip_prob', 'fivecrop', 'hflip_logit', 'temperature', 'amp']

def scalar_metrics(y, probs):
    m = compute_metrics(np.asarray(y), probs.argmax(1), probs)
    return {k: m[k] for k in ('macro_f1', 'top1', 'ece', 'balanced_acc', 'nll')}

def eda(base, output):
    import matplotlib.pyplot as plt
    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    frames = dataset.load_split(base.labels_dir)
    checks = dataset.check_split(*frames, base.images_dir)
    write_json(output/'split_checks.json', checks)
    counts = pd.DataFrame(checks['per_class'], index=dataset.CLASS_NAMES)
    counts.to_csv(output/'class_counts.csv')
    counts.plot.bar(figsize=(12, 5)); plt.ylabel('Images'); plt.tight_layout()
    plt.savefig(output/'class_distribution.png', dpi=160); plt.close()
    all_df = pd.concat(frames)
    sizes = {}
    for name in all_df.Filename:
        with Image.open(Path(base.images_dir)/name) as im:
            key = f'{im.size}/{im.mode}'; sizes[key] = sizes.get(key, 0) + 1
    write_json(output/'image_sizes.json', sizes)
    fig, axes = plt.subplots(9, 3, figsize=(9, 24))
    for label in range(9):
        for j, name in enumerate(all_df[all_df.Label == label].Filename.iloc[:3]):
            with Image.open(Path(base.images_dir)/name) as im: axes[label, j].imshow(im.convert('RGB'))
            axes[label, j].axis('off'); axes[label, j].set_title(dataset.CLASS_NAMES[label])
    fig.tight_layout(); fig.savefig(output/'samples.png', dpi=100); plt.close(fig)
    # Visual inspection of train augmentation, without using test pixels to tune normalization.
    loader = dataset.make_loader(frames[0], base.images_dir, dataset.build_transforms(True), 8, False, num_workers=0)
    x, y, _ = next(iter(loader))
    fig, axes = plt.subplots(2, 4, figsize=(12, 6))
    for ax, image, label in zip(axes.flat, x, y):
        im = image.permute(1, 2, 0).numpy() * dataset.IMAGENET_STD + dataset.IMAGENET_MEAN
        ax.imshow(np.clip(im, 0, 1)); ax.set_title(dataset.CLASS_NAMES[int(label)]); ax.axis('off')
    fig.tight_layout(); fig.savefig(output/'augmentations.png', dpi=130); plt.close(fig)
    print(counts, '\nSplit checks:', checks)
    return checks

def check_pipeline(base, output):
    # Actual backbone overfits a small batch; no test loader is constructed here.
    from .losses import FocalLoss, mix_batch
    set_seed(base.seed)
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    df = dataset.load_split(base.labels_dir)[0].groupby('Label').head(1).iloc[:4]
    loader = dataset.make_loader(df, base.images_dir, dataset.build_transforms(False, 64), 4, False, num_workers=0)
    x, y, _ = next(iter(loader)); x, y = x.to(device), y.to(device)
    model = models.build_model('resnet18', pretrained=False, init='scratch').to(device).eval()
    initial = float(torch.nn.functional.cross_entropy(model(x), y).detach())
    opt = torch.optim.Adam(model.parameters(), lr=.003)
    loss_history = []
    for _ in range(150):
        opt.zero_grad(); z = model(x); loss = torch.nn.functional.cross_entropy(z, y)
        loss.backward(); opt.step(); loss_history.append(float(loss.detach()))
        if loss_history[-1] < .03: break
    z = torch.randn(8, 9, device=device); target = torch.arange(8, device=device)
    error = float((FocalLoss(0)(z, target) - torch.nn.functional.cross_entropy(z, target)).abs())
    mixed, (_, _, lam) = mix_batch(x, y)
    evidence = dict(initial_ce=initial, expected_ce=float(np.log(9)), overfit_final_ce=loss_history[-1],
                    overfit_steps=len(loss_history), loss_history=loss_history, focal_gamma0_error=error,
                    cutmix_lambda=lam, train_samples=df.Filename.tolist(), backbone='resnet18 scratch, eval mode, 64px')
    write_json(Path(output)/'pipeline_checks.json', evidence)
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(1, len(mixed), figsize=(12, 3))
    for ax, image in zip(axes, mixed.detach().cpu()):
        pixels = image.permute(1, 2, 0).numpy() * dataset.IMAGENET_STD + dataset.IMAGENET_MEAN
        ax.imshow(np.clip(pixels, 0, 1)); ax.axis('off')
    fig.suptitle(f'CutMix lambda={lam:.4f}'); fig.tight_layout()
    fig.savefig(Path(output)/'cutmix.png', dpi=130); plt.close(fig)
    if error > 1e-6 or loss_history[-1] >= .1:
        raise RuntimeError(f'Pipeline check failed. Inspect {output}/pipeline_checks.json')
    del model, opt; gc.collect()
    if device.type == 'cuda': torch.cuda.empty_cache()
    print('Pipeline checks passed:', {k:v for k,v in evidence.items() if k != 'loss_history'})
    return evidence

def load_model(cfg):
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    model = models.build_model(cfg.backbone, pretrained=False, drop_rate=cfg.drop_rate, init=cfg.init).to(device)
    model.load_state_dict(torch.load(run_dir(cfg)/'best.pt', map_location=device, weights_only=False)['model'])
    return model.eval(), device

def views(x, method, img_size):
    if method == 'fivecrop': return inf.views_multicrop(x, img_size)
    if method.startswith('hflip'): return [x, inf.view_hflip(x)]
    return [x]

def forward_method(model, x, method, img_size, temperature=1.):
    with torch.autocast(device_type=x.device.type, enabled=method == 'amp' and x.device.type == 'cuda'):
        zs = [model(v).float() for v in views(x, method, img_size)]
    if method == 'hflip_logit':
        logits = torch.stack(zs).mean(0)
    elif len(zs) > 1:
        logits = torch.stack([z.softmax(1) for z in zs]).mean(0).clamp_min(1e-12).log()
    else:
        logits = zs[0]
    return logits, (logits / temperature).softmax(1)

def predict_method(model, cfg, split, method, temperature=1.):
    device = next(model.parameters()).device
    frame = dataset.load_split(cfg.labels_dir, cfg.fold)[{'train':0, 'val':1, 'test':2}[split]]
    env = json.loads((run_dir(cfg)/'environment.json').read_text())['normalization']
    if method == 'fivecrop':
        transform = T.Compose([T.Resize((round(cfg.img_size/.875),)*2), T.ToTensor(), T.Normalize(env['mean'], env['std'])])
    else:
        transform = dataset.build_transforms(False, cfg.img_size, mean=env['mean'], std=env['std'])
    loader = dataset.make_loader(frame, cfg.images_dir, transform, cfg.batch_size, False,
                                 num_workers=cfg.num_workers, seed=cfg.seed)
    names, labels, logits, probs = [], [], [], []
    model.eval()
    with torch.inference_mode():
        for x, y, files in loader:
            z, p = forward_method(model, x.to(device), method, cfg.img_size, temperature)
            names.extend(files); labels.extend(y.tolist())
            logits.append(z.cpu().numpy()); probs.append(p.cpu().numpy())
    return names, np.asarray(labels), np.concatenate(logits), np.concatenate(probs)

def measure_method(model, cfg, method, batch=1, temperature=1.):
    device = next(model.parameters()).device
    side = round(cfg.img_size/.875) if method == 'fivecrop' else cfg.img_size
    x = torch.randn(batch, 3, side, side, device=device)
    with torch.inference_mode():
        result = benchmark.bench(lambda: forward_method(model, x, method, cfg.img_size, temperature),
                                 sync=torch.cuda.synchronize if device.type == 'cuda' else None)
    result.update(batch=batch, gpu=torch.cuda.get_device_name() if device.type == 'cuda' else 'CPU',
                  dtype='amp' if method == 'amp' else 'fp32', img_size=cfg.img_size,
                  method=method, images_per_s=batch*1000/result['p50'], torch=torch.__version__,
                  preprocessing='image decode/resize/normalize excluded; TTA views + aggregate included')
    return result

def inference_experiments(cfg, root):
    root = Path(root); dest = root/'inference'; dest.mkdir(exist_ok=True)
    model, device = load_model(cfg)
    rows, latency = [], []
    for i, method in enumerate(METHODS):
        if method == 'amp' and device.type != 'cuda': continue
        names, y, z, probs = predict_method(model, cfg, 'val', method)
        temperature = inf.fit_temperature(z, y) if method == 'temperature' else 1.
        probs = inf.apply_temperature(z, temperature)
        np.savez_compressed(dest/f'I{i:02}_val.npz', filenames=np.asarray(names), y=y, logits=z, probs=probs)
        measurements = [measure_method(model, cfg, method, b, temperature) for b in (1, 32)]
        latency.extend([{'exp_id': f'I{i:02}', **r} for r in measurements])
        rows.append({'exp_id':f'I{i:02}', 'method':method, 'checkpoint':str(run_dir(cfg)/'best.pt'),
                     'K':5 if method == 'fivecrop' else 2 if method.startswith('hflip') else 1,
                     'temperature':temperature, **scalar_metrics(y, probs), **measurements[0]})
        print('Inference:', method, rows[-1]['macro_f1'], rows[-1]['p95'])
    baseline_ms = rows[0]['p50']
    for row in rows: row['relative_cost'] = row['p50']/baseline_ms
    # Quality-first selection, then calibration and latency to break ties.
    winner = sorted(rows, key=lambda r: (-r['macro_f1'], r['ece'], r['p95']))[0]
    write_json(dest/'rows.json', rows); write_json(dest/'latency.json', latency)
    del model; gc.collect()
    if device.type == 'cuda': torch.cuda.empty_cache()
    return winner

def screen(base, root, backbones=None):
    root = Path(root); root.mkdir(parents=True, exist_ok=True)
    if (root/'final_lock.json').exists():
        raise RuntimeError('Final configuration already locked. Run final/export; do not rescreen after test.')
    backbones = backbones or models.SUGGESTED_BACKBONES
    results = [run(replace(base, exp_id=f'B{i:02}', backbone=name)) for i, name in enumerate(backbones, 1)]
    # Add comparable baseline forward latencies for every backbone.
    for result in results:
        cfg = Config(**{k:result[k] for k in asdict(base)})
        model, device = load_model(cfg)
        result['latency_batch1'] = benchmark.latency_report(model, 1, cfg.img_size, device=str(device))['p50']
        write_json(run_dir(cfg)/'result.json', result)
        del model; gc.collect()
        if device.type == 'cuda': torch.cuda.empty_cache()
    winner = sorted(results, key=lambda r: (-r['macro_f1'], r['latency_batch1']))[0]
    baseline = replace(base, backbone=winner['backbone'], exp_id='T00')
    # Initialization, augmentation and loss: only one factor changes from T00.
    changes = [({}, 'baseline'), ({'init':'scratch'}, 'A/init'), ({'init':'frozen'}, 'A/init'),
               ({'mix':'cutmix'}, 'B/mix'), ({'aug':'color'}, 'B/aug'),
               ({'loss':'ls', 'label_smoothing':.1}, 'C/loss'),
               ({'loss':'focal'}, 'C/loss'), ({'loss':'ce_weighted'}, 'C/loss')]
    training = []
    for i, (delta, axis) in enumerate(changes):
        result = run(replace(baseline, exp_id=f'T{i:02}', **delta))
        result.update(axis=axis, changed=json.dumps(delta), delta_f1=result['macro_f1']-training[0]['macro_f1'] if training else 0.)
        training.append(result)
    # Combine the strongest individually tested value per independent factor.
    combination = {}
    for keys in [('init',), ('mix', 'aug'), ('loss', 'label_smoothing')]:
        candidates = [r for r in training if any(k in json.loads(r['changed']) for k in keys)]
        best = max(candidates, key=lambda r:r['macro_f1'])
        if best['macro_f1'] > training[0]['macro_f1']:
            combination.update(json.loads(best['changed']))
    combo = run(replace(baseline, exp_id='T08', **combination))
    combo.update(axis='combination', changed=json.dumps(combination), delta_f1=combo['macro_f1']-training[0]['macro_f1'])
    training.append(combo)
    chosen = max(training, key=lambda r:r['macro_f1'])
    final_cfg = Config(**{k:chosen[k] for k in asdict(base)})
    method = inference_experiments(final_cfg, root)
    plan = dict(baseline=asdict(baseline), chosen=asdict(final_cfg), method=method['method'],
                selection='validation macro-F1; inference ties broken by ECE then p95',
                backbone_rows=results, training_rows=training)
    write_json(root/'selection.json', plan)
    print('Selected on validation:', chosen['exp_id'], chosen['backbone'], method['method'])
    return plan

def finalize(root, seeds=(0, 1, 2)):
    root = Path(root)
    if len(seeds) < 3 or len(set(seeds)) != len(seeds): raise ValueError('At least 3 distinct seeds required')
    selection = json.loads((root/'selection.json').read_text())
    lock_path = root/'final_lock.json'
    lock = {'baseline':selection['baseline'], 'chosen':selection['chosen'],
            'method':selection['method'], 'seeds':list(seeds)}
    if lock_path.exists() and json.loads(lock_path.read_text()) != lock:
        raise RuntimeError('Locked final configuration cannot be changed')
    write_json(lock_path, lock)
    rows, per_class_rows, latency_rows = [], [], []
    for group, source, method in [('T00', lock['baseline'], 'single'), ('F01', lock['chosen'], lock['method'])]:
        for seed in seeds:
            cfg = replace(Config(**source), exp_id=group, seed=seed)
            result = run(cfg)
            folder = run_dir(cfg); final_file = folder/'final_result.json'
            pred_file = Path(cfg.pred_dir)/f'{group}_seed{seed}_test.csv'
            if final_file.exists():
                record = json.loads(final_file.read_text())
            else:
                started = folder/'test_started.json'
                if started.exists():
                    raise RuntimeError(f'{folder}: previous test interrupted. Preserve evidence; do not silently rerun test.')
                model, device = load_model(cfg)
                names, y, z, _ = predict_method(model, cfg, 'val', method)
                # Fit a separate T per seed on its validation output; fixed rule locked before test.
                temperature = inf.fit_temperature(z, y) if group == 'F01' else 1.
                val_metrics = scalar_metrics(y, inf.apply_temperature(z, temperature))
                latency = [measure_method(model, cfg, method, batch, temperature) for batch in (1, 32)]
                write_json(folder/'test_started.json', {'group':group, 'seed':seed, 'method':method, 'T':temperature})
                names, y, z, probs = predict_method(model, cfg, 'test', method, temperature)
                before = scalar_metrics(y, inf.softmax(z))
                save_predictions(pred_file, names, y, probs)
                np.savez_compressed(folder/'test_logits.npz', filenames=np.asarray(names), y=y, logits=z, probs=probs)
                metrics = compute_metrics(y, probs.argmax(1), probs)
                pc = [dict(group=group, seed=seed, label=i, class_name=dataset.CLASS_NAMES[i],
                           **{k:float(metrics[k][i]) for k in ('support','precision','recall','f1')}) for i in range(9)]
                record = dict(exp_id=group, seed=seed, backbone=cfg.backbone, method=method,
                              temperature=temperature, macro_f1_val=val_metrics['macro_f1'],
                              **scalar_metrics(y, probs), ece_before=before['ece'],
                              per_class=pc, latency=latency)
                write_json(final_file, record)
                del model; gc.collect()
                if device.type == 'cuda': torch.cuda.empty_cache()
            rows.append({k:v for k,v in record.items() if k not in ('per_class','latency')})
            per_class_rows.extend(record['per_class'])
            latency_rows.extend([dict(exp_id=group, seed=seed, **r) for r in record['latency']])
    write_json(root/'final_rows.json', rows); write_json(root/'per_class_rows.json', per_class_rows)
    write_json(root/'final_latency.json', latency_rows)
    # Official tool validates all filenames/labels and produces comparable aggregates.
    import subprocess, sys
    base = Config(**lock['baseline']); evaluator = Path(__file__).resolve().parents[1]/'eval.py'
    if not evaluator.exists(): evaluator = Path(__file__).resolve().parents[2]/'eval.py'
    common = ['--test-csv', str(Path(base.labels_dir)/'test_subset0.csv'), '--labels', str(Path(base.labels_dir)/'labels.csv')]
    for group in ('T00','F01'):
        subprocess.run([sys.executable, '-X', 'utf8', str(evaluator), 'score', '--pred', str(Path(base.pred_dir)/f'{group}_seed*_test.csv'),
                        *common, '--tag', group, '--out', str(root/'eval_out')], check=True)
    with (root/'grade.txt').open('w', encoding='utf-8') as output:
        subprocess.run([sys.executable, '-X', 'utf8', str(evaluator), 'grade', '--final', str(Path(base.pred_dir)/'F01_seed*_test.csv'),
                        '--baseline', str(Path(base.pred_dir)/'T00_seed*_test.csv'), *common], check=True, stdout=output)
    return rows

def export(root):
    import matplotlib.pyplot as plt
    root = Path(root)
    plan = json.loads((root/'selection.json').read_text())
    def read(name, default):
        p = root/name; return json.loads(p.read_text()) if p.exists() else default
    final = pd.DataFrame(read('final_rows.json', []))
    inference = pd.DataFrame(read('inference/rows.json', []))
    latency = pd.DataFrame(read('inference/latency.json', []) + read('final_latency.json', []))
    training = pd.DataFrame(plan['training_rows']); backbones = pd.DataFrame(plan['backbone_rows'])
    summary = pd.concat([backbones.assign(stage='backbone'), training.assign(stage='training')], ignore_index=True)
    summary = summary.sort_values('macro_f1', ascending=False).head(10)
    sheets = dict(Backbones=backbones, Training=training, Inference=inference, Final=final,
                  PerClass=pd.DataFrame(read('per_class_rows.json', [])), Latency=latency, Summary=summary)
    if not final.empty:
        numeric = ['macro_f1', 'top1', 'ece', 'macro_f1_val']
        aggregate = final.groupby('exp_id')[numeric].agg(['mean','std'])
        aggregate.columns = ['_'.join(c) for c in aggregate.columns]
        sheets['FinalSummary'] = aggregate.reset_index()
        pc_summary = sheets['PerClass'].groupby(['group', 'class_name'])[['precision','recall','f1']].agg(['mean','std'])
        pc_summary.columns = ['_'.join(c) for c in pc_summary.columns]
        sheets['PerClassSummary'] = pc_summary.reset_index()
    with pd.ExcelWriter(root/'results.xlsx', engine='openpyxl') as writer:
        for name, frame in sheets.items():
            # Nested metadata is serialized rather than becoming invalid Excel cells.
            frame = frame.map(lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else x)
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]; ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='244062')
            from openpyxl.utils import get_column_letter
            for i in range(1, ws.max_column+1): ws.column_dimensions[get_column_letter(i)].width = 20
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.scatter(inference.p95, inference.macro_f1)
    for _, row in inference.iterrows(): ax.annotate(row.method, (row.p95, row.macro_f1), fontsize=8)
    ax.set(xlabel='p95 batch 1 (ms)', ylabel='Macro-F1 validation'); fig.tight_layout()
    fig.savefig(root/'curves'/'inference_tradeoff.png', dpi=160); plt.close(fig)
    error_lines = []
    if not final.empty:
        base = Config(**plan['baseline'])
        for group in ('T00', 'F01'):
            files = sorted(Path(base.pred_dir).glob(f'{group}_seed*_test.csv'))
            cms = []
            for file in files:
                df = pd.read_csv(file); probs = df[[f'p{i}' for i in range(9)]].to_numpy()
                cms.append(compute_metrics(df.y_true.to_numpy(), df.y_pred.to_numpy(), probs)['confusion'])
                errors = df[df.y_true != df.y_pred].copy()
                errors['confidence'] = probs.max(1)[df.y_true != df.y_pred]
                errors.sort_values('confidence', ascending=False).to_csv(root/f'{file.stem}_errors.csv', index=False)
            cm = np.mean(cms, axis=0)
            fig, ax = plt.subplots(figsize=(10, 8)); image = ax.imshow(cm, cmap='Blues')
            ax.set_xticks(range(9), dataset.CLASS_NAMES, rotation=90); ax.set_yticks(range(9), dataset.CLASS_NAMES)
            ax.set(xlabel='Predicted', ylabel='True', title=f'{group}: mean counts across seeds')
            for i in range(9):
                for j in range(9): ax.text(j, i, f'{cm[i,j]:.1f}', ha='center', va='center', fontsize=7)
            fig.colorbar(image, ax=ax); fig.tight_layout(); fig.savefig(root/'curves'/f'{group}_confusion.png', dpi=160); plt.close(fig)
            error_lines.append(f'- {group}: Chinee Apple -> Snake Weed = {cm[0,7]:.1f}; chiều ngược lại = {cm[7,0]:.1f} ảnh trung bình/seed.')
        # Visual errors for the first final seed, not cherry-picked best seed.
        first = sorted(Path(base.pred_dir).glob('F01_seed*_test.csv'))[0]
        df = pd.read_csv(first); mistakes = df[df.y_true != df.y_pred].head(12)
        if len(mistakes):
            fig, axes = plt.subplots(3, 4, figsize=(14, 11))
            for ax in axes.flat: ax.axis('off')
            for ax, row in zip(axes.flat, mistakes.itertuples()):
                with Image.open(Path(base.images_dir)/row.Filename) as im: ax.imshow(im.convert('RGB'))
                ax.set_title(f'True: {dataset.CLASS_NAMES[row.y_true]}\nPred: {dataset.CLASS_NAMES[row.y_pred]}', fontsize=8)
            fig.tight_layout(); fig.savefig(root/'curves'/'misclassified.png', dpi=140); plt.close(fig)
    report = ['# Báo cáo DeepWeeds — kết quả chạy thật', '',
              '> Bản tổng hợp tự động. Bổ sung diễn giải ảnh lỗi, giả thuyết và hạn chế trước khi nộp.', '',
              '## Dữ liệu và thiết lập',
              'Fold 0 gốc; train/val/test tách biệt. Chọn cấu hình chỉ bằng validation macro-F1. '
              'Chia ngẫu nhiên không theo địa điểm có thể làm kết quả lạc quan.',
              f'Cấu hình baseline: `{json.dumps(plan["baseline"], ensure_ascii=False)}`',
              '## Backbone', backbones[['exp_id','backbone','macro_f1','top1','params_m','gmac','train_seconds']].to_markdown(index=False),
              '## Huấn luyện', training[['exp_id','axis','changed','macro_f1','delta_f1']].to_markdown(index=False),
              'Sàng lọc dùng một seed; chưa coi chênh lệch nhỏ là bằng chứng chắc chắn.',
              '## Suy luận', inference[['method','macro_f1','top1','ece','p95','relative_cost']].to_markdown(index=False),
              '![Đánh đổi](curves/inference_tradeoff.png)',
              f'Cấu hình được chọn: {plan["chosen"]["exp_id"]}, {plan["chosen"]["backbone"]}; suy luận {plan["method"]}.',
              '## Chung kết']
    if final.empty:
        report.append('Chưa chạy test chung kết. Không có số liệu test để kết luận.')
    else:
        report.append(sheets['FinalSummary'].to_markdown(index=False))
        stats = final.groupby('exp_id').macro_f1.agg(['mean','std'])
        delta = stats.loc['F01','mean'] - stats.loc['T00','mean']; noise = stats['std'].max()
        report.append(f'Δ macro-F1 test = {delta:.6f}; std lớn nhất = {noise:.6f}. '
                      + ('Cải thiện vượt std.' if delta > noise else 'Chưa có bằng chứng cải thiện vượt nhiễu.'))
        f = final[final.exp_id == 'F01']; gap = (f.macro_f1_val-f.macro_f1).abs().mean()
        report.append(f'Chênh lệch tuyệt đối val/test trung bình: {gap:.6f}. ECE trước/sau: '
                      f'{f.ece_before.mean():.6f}/{f.ece.mean():.6f}.')
        realtime = latency[(latency.exp_id == 'F01') & (latency.batch == 1)]
        report.append(f'p95 chung kết batch 1 trung bình: {realtime.p95.mean():.2f} ms. '
                      + ('Đạt ngân sách 100 ms trên phần cứng đã đo.' if realtime.p95.max() <= 100 else 'Chưa đạt ổn định ngân sách 100 ms.'))
        report.extend(error_lines + ['![Ma trận nhầm lẫn](curves/F01_confusion.png)', '![Ảnh lỗi](curves/misclassified.png)'])
    report += ['## Kết luận và hạn chế',
               'Đối chiếu results.xlsx để giải thích yếu tố đóng góp nhiều nhất, cộng dồn/triệt tiêu và lựa chọn cho robot. '
               'Không suy diễn FLOPs thành độ trễ. Khác biệt kiến trúc còn chịu ảnh hưởng của bộ trọng số tiền huấn luyện.',
               'Chỉ một fold; chưa kiểm chứng địa điểm/mùa mới. Tái lập dùng seed cố định nhưng kernel GPU có thể vẫn không hoàn toàn xác định. '
               'Loss train là objective đang thử, loss val luôn CE; không so trực tiếp trị số của hai loại loss khác nhau.',
               '## Phụ lục', 'Cấu hình/log/checkpoint: runs/<exp_id>/seed<k>. Chi tiết chấm chính thức: eval_out/ và grade.txt.']
    (root/'report.md').write_text('\n\n'.join(report), encoding='utf-8')
    # Submission archive excludes dataset and checkpoints; scripts and evidence are included.
    archive = root/'submission'; archive.mkdir(exist_ok=True)
    for name in ('results.xlsx','report.md','grade.txt','final_lock.json','selection.json'):
        if (root/name).exists(): shutil.copy2(root/name, archive/name)
    for name in ('curves','predictions','eda','eval_out'):
        if (root/name).exists(): shutil.copytree(root/name, archive/name, dirs_exist_ok=True)
    code_root = Path(__file__).parent
    shutil.copytree(code_root, archive/'code'/'lab_code', dirs_exist_ok=True, ignore=shutil.ignore_patterns('__pycache__'))
    evaluator = code_root.parent/'eval.py'
    if not evaluator.exists(): evaluator = code_root.parents[1]/'eval.py'
    shutil.copy2(evaluator, archive/'code'/'eval.py')
    for file in (root/'runs').rglob('history.csv'):
        dest = archive/'logs'/file.relative_to(root); dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(file,dest)
    for name in ('config.json','environment.json','final_result.json'):
        for file in (root/'runs').rglob(name):
            dest = archive/'logs'/file.relative_to(root); dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(file,dest)
    (archive/'README.md').write_text('''# DeepWeeds lab submission
Set GPU and Internet on Kaggle, attach DeepWeeds (images + original fold-0 CSVs), import deepweeds_kaggle.ipynb.
Run setup, EDA, pipeline checks, screen, then final and export in order.
Code: python -m lab_code.suite --stage screen --output /kaggle/working/deepweeds_lab
Then --stage final and --stage export with the same paths/epochs.
See logs/**/environment.json for versions, GPU, normalization and weight tags.
Checkpoint links and your public Kaggle notebook URL: add before submission.
The generated report requires your interpretation of error images and hypotheses.
''', encoding='utf-8')
    notebook = code_root.parent/'deepweeds_kaggle.ipynb'
    if notebook.exists(): shutil.copy2(notebook, archive/'code'/notebook.name)
    shutil.make_archive(str(root/'deepweeds_submission'), 'zip', archive)
    print('Created:', root/'results.xlsx', root/'report.md', root/'deepweeds_submission.zip')

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--stage', choices=['prepare','screen','final','export','all'], default='prepare')
    parser.add_argument('--input-root', default='/kaggle/input')
    parser.add_argument('--images-dir'); parser.add_argument('--labels-dir')
    parser.add_argument('--output', default='/kaggle/working/deepweeds_lab')
    parser.add_argument('--epochs', type=int, default=12); parser.add_argument('--batch-size', type=int, default=32)
    args = parser.parse_args()
    root = Path(args.output); root.mkdir(parents=True, exist_ok=True)
    images, labels = dataset.discover(args.input_root, args.images_dir, args.labels_dir)
    base = Config(images_dir=images, labels_dir=labels, epochs=args.epochs, batch_size=args.batch_size,
                  out_dir=str(root/'runs'), pred_dir=str(root/'predictions'), curves_dir=str(root/'curves'))
    if args.stage in ('prepare','all'): eda(base, root/'eda'); check_pipeline(base, root/'eda')
    if args.stage in ('screen','all'): screen(base, root)
    if args.stage in ('final','all'): finalize(root)
    if args.stage in ('export','all'): export(root)

if __name__ == '__main__': main()
